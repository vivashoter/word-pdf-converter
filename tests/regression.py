import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz
from docx import Document

from app import (
    convert_la350_to_docx,
    is_la350_form,
    pdf_widget_count,
)

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures"
OUT = ROOT / "test-output"
OUT.mkdir(exist_ok=True)


def fail(message):
    raise AssertionError(message)


def check_docx(path):
    if not path.exists() or path.stat().st_size == 0:
        fail("Generated DOCX is missing or empty.")

    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text

    required = [
        "Notice of Available Language",
        "Assistance—Service Provider",
        "Information about the services provided",
        "Services",
        "Languages Available",
        "Types of Language",
        "Service Area",
        "Date:",
        "LA-350, Page 1 of 1",
    ]
    for phrase in required:
        if phrase not in text:
            fail(f"Missing expected LA-350 content: {phrase}")


def render_docx(docx_path):
    pdf_dir = OUT / "rendered"
    pdf_dir.mkdir(exist_ok=True)
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(pdf_dir),
            str(docx_path),
        ],
        check=True,
        timeout=120,
    )
    pdf_path = pdf_dir / (docx_path.stem + ".pdf")
    if not pdf_path.exists():
        fail("LibreOffice did not render the generated DOCX.")
    return pdf_path


def check_rendered_pdf(pdf_path):
    pdf = fitz.open(pdf_path)
    try:
        if pdf.page_count != 1:
            fail(f"LA-350 regression: expected 1 page, got {pdf.page_count}.")

        page = pdf[0]
        text = page.get_text("text") or ""
        if "LA-350" not in text:
            fail("Rendered page is missing LA-350 text.")

        # Catch accidental page-size/layout regressions.
        width = page.rect.width
        height = page.rect.height
        if not (600 <= width <= 620 and 780 <= height <= 800):
            fail(f"Unexpected rendered page size: {width:.1f} x {height:.1f} pt.")

        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        pix.save(str(OUT / "la350-render.png"))
    finally:
        pdf.close()


def main():
    fixture = FIXTURES / "LA-350.pdf"

    if not fixture.exists():
        print("LA-350 fixture is not in the repository yet.")
        print("Core import test passed; add tests/fixtures/LA-350.pdf to enable full visual regression testing.")
        return

    count = pdf_widget_count(str(fixture))
    if count <= 0:
        fail("LA-350 fixture should contain interactive form widgets.")

    if not is_la350_form(str(fixture)):
        fail("LA-350 detector did not recognize the fixture.")

    output_docx = OUT / "LA-350-test.docx"
    convert_la350_to_docx(str(fixture), str(output_docx))
    check_docx(output_docx)

    rendered_pdf = render_docx(output_docx)
    check_rendered_pdf(rendered_pdf)

    shutil.copy2(rendered_pdf, OUT / "LA-350-test.pdf")
    print("LA-350 regression test passed.")


if __name__ == "__main__":
    main()
